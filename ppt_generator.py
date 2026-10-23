import requests
import re
import time
import os
from html import unescape
try:
    import win32com.client as win32
except Exception:
    win32 = None

try:
    import tkinter as tk
    from tkinter import ttk, filedialog
except Exception:
    tk = None
    ttk = None
    filedialog = None

try:
    import google.generativeai as genai
    _GENAI_AVAILABLE = True
except Exception:
    genai = None
    _GENAI_AVAILABLE = False

import json

# Set your Gemini API key here (or export environment variable GEMINI_API_KEY)
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "")
GROQ_API_KEY = ""
try:
    import config as cfg
    if not GEMINI_API_KEY:
        GEMINI_API_KEY = cfg.CONFIG_MAP.get("gemini_api_key", "")
    GROQ_API_KEY = cfg.CONFIG_MAP.get("groq_api_key", "")
except Exception:
    pass

DEFAULT_GEMINI_MODEL = "models/gemini-2.5-flash"

def _clean_text(t):
    return re.sub(r"\s+", " ", unescape(re.sub(r"<.*?>", " ", t))).strip()

def search_google_links(query, max_results=10):
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/118.0.0.0 Safari/537.36",
            "Accept-Language": "en-US,en;q=0.9",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"
        }
        url = f"https://www.google.com/search?q={requests.utils.quote(query)}&hl=en&num={max_results}"
        r = requests.get(url, headers=headers, timeout=10)
        if r.status_code != 200:
            return []
        html = r.text
        links = []
        for m in re.finditer(r"href=\"/url\?q=(.*?)&amp;", html):
            link = unescape(m.group(1))
            if not link.startswith("http"):
                continue
            if "google." in link:
                continue
            links.append(link)
            if len(links) >= max_results:
                break
        titles = []
        for m in re.finditer(r"<h3[^>]*>(.*?)</h3>", html):
            titles.append(_clean_text(m.group(1)))
        items = []
        for i, link in enumerate(links):
            title = titles[i] if i < len(titles) else link
            items.append((title, link))
        return items
    except Exception:
        return []

def fetch_page_summary(url, max_chars=800):
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/118.0.0.0 Safari/537.36",
            "Accept-Language": "en-US,en;q=0.9",
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"
        }
        r = requests.get(url, headers=headers, timeout=10)
        if r.status_code != 200:
            return ""
        html = r.text
        html = re.sub(r"<script[\s\S]*?</script>", " ", html, flags=re.I)
        html = re.sub(r"<style[\s\S]*?</style>", " ", html, flags=re.I)
        # meta description
        md = re.search(r"<meta[^>]*name=\"description\"[^>]*content=\"([^\"]+)\"", html, flags=re.I)
        meta_desc = _clean_text(md.group(1)) if md else ""
        heads = re.findall(r"<(h1|h2|h3)[^>]*>([\s\S]*?)</\1>", html, flags=re.I)
        paras = re.findall(r"<p[^>]*>([\s\S]*?)</p>", html, flags=re.I)
        lis = re.findall(r"<li[^>]*>([\s\S]*?)</li>", html, flags=re.I)
        buf = []
        if meta_desc:
            buf.append(meta_desc)
        for _, h in heads[:5]:
            ht = _clean_text(h)
            if ht:
                buf.append(ht)
        for p in paras[:12]:
            text = _clean_text(p)
            if len(text) >= 40:
                buf.append(text)
            if sum(len(x) for x in buf) >= max_chars:
                break
        if len(buf) < 3:
            for li in lis[:10]:
                lt = _clean_text(li)
                if lt:
                    buf.append(lt)
        out = "\n".join(buf).strip()
        return out[:max_chars]
    except Exception:
        return ""

def format_popup(slide_count):
    root = tk.Tk()
    root.title("JARVIS — Slide Format")
    root.geometry("520x" + str(160 + slide_count * 40))
    style = ttk.Style()
    style.theme_use('clam')
    ttk.Label(root, text="Enter titles for each slide").pack(pady=10)
    entries = []
    for i in range(slide_count):
        frame = ttk.Frame(root)
        frame.pack(pady=4)
        ttk.Label(frame, text=f"Slide {i+1}").pack(side=tk.LEFT, padx=6)
        e = ttk.Entry(frame, width=48)
        e.pack(side=tk.LEFT)
        entries.append(e)
    res = {"ok": False, "titles": []}
    def on_ok():
        res["ok"] = True
        res["titles"] = [x.get().strip() or f"Slide {i+1}" for i, x in enumerate(entries)]
        root.destroy()
    def on_cancel():
        res["ok"] = False
        root.destroy()
    btns = ttk.Frame(root)
    btns.pack(pady=10)
    ttk.Button(btns, text="OK", command=on_ok).pack(side=tk.LEFT, padx=8)
    ttk.Button(btns, text="Cancel", command=on_cancel).pack(side=tk.LEFT, padx=8)
    root.mainloop()
    return res["titles"] if res["ok"] else None

def select_existing_ppt():
    root = tk.Tk()
    root.withdraw()
    path = filedialog.askopenfilename(title="Select themed PPT", filetypes=[("PowerPoint", "*.pptx")])
    try:
        root.destroy()
    except Exception:
        pass
    return path or None

def choose_theme_with_gemini(topic):
    try:
        from providers.provider_manager import ProviderManager
        prompt = (
            f"Create a concise JSON theme for a presentation about '{topic}'. Include keys: background_gradient(start,end hex), text(primary,secondary hex), fonts(title,body names). Only output valid JSON."
        )
        text = ProviderManager.gemini(prompt)
        import json
        return json.loads(text)
    except Exception:
        return None

def _hex_to_rgb(hex_color):
    try:
        h = hex_color.strip().lstrip('#')
        return int(h[4:6] + h[2:4] + h[0:2], 16)
    except Exception:
        return int('FFFFFF', 16)

def _apply_theme(pres, slide, theme):
    try:
        if not theme:
            theme = {
                "background_gradient": {"start": "#07090D", "end": "#0A1F30"},
                "text": {"primary": "#00D2FF", "secondary": "#8DA4B4"},
                "fonts": {"title": "Outfit", "body": "Outfit"},
                "accent": "#00FFCC"
            }
        else:
            if "background_gradient" not in theme:
                theme["background_gradient"] = {"start": "#07090D", "end": "#0A1F30"}
            if "text" not in theme:
                theme["text"] = {"primary": "#00D2FF", "secondary": "#8DA4B4"}
            if "fonts" not in theme:
                theme["fonts"] = {"title": "Outfit", "body": "Outfit"}
            if "accent" not in theme:
                theme["accent"] = theme["text"].get("primary", "#00FFCC")

        slide.FollowMasterBackground = False
        fill = slide.Background.Fill
        fill.Visible = True
        try:
            start_color = _hex_to_rgb(theme["background_gradient"]["start"])
            end_color = _hex_to_rgb(theme["background_gradient"]["end"])
            fill.TwoColorGradient(1, 1) 
            fill.ForeColor.RGB = start_color
            fill.BackColor.RGB = end_color
        except Exception:
            try:
                fill.Solid()
                fill.ForeColor.RGB = _hex_to_rgb(theme["background_gradient"]["start"])
            except Exception:
                pass
            
        primary_color = _hex_to_rgb(theme["text"]["primary"])
        secondary_color = _hex_to_rgb(theme["text"]["secondary"])
        accent_color = _hex_to_rgb(theme["accent"])
        title_font = theme["fonts"]["title"]
        body_font = theme["fonts"]["body"]
        
        shapes_to_format = list(slide.Shapes)
        for shp in shapes_to_format:
            # Skip custom card shapes or custom layout shapes
            if shp.Type == 5 or (shp.Name and any(k in shp.Name for k in ["Card", "Col", "Bullet", "Subtitle"])):
                continue
                
            if shp.HasTextFrame:
                tf = shp.TextFrame
                if tf.HasText:
                    is_title = False
                    try:
                        if shp.Name and "Title" in shp.Name:
                            is_title = True
                        elif shp.Type == 14: 
                            if shp.PlaceholderFormat.Type == 1: 
                                is_title = True
                    except Exception:
                        pass
                    
                    if shp.Top < 120:
                        is_title = True
                        
                    tr = tf.TextRange
                    tr.Font.Name = title_font if is_title else body_font
                    tr.Font.Color.RGB = primary_color if is_title else secondary_color
                    if is_title:
                        tr.Font.Bold = True
                        tr.Font.Size = 36
                        
                        try:
                            left = shp.Left
                            top = shp.Top + shp.Height + 5
                            width = shp.Width
                            divider = slide.Shapes.AddShape(1, left, top, width, 2) 
                            divider.Fill.Solid()
                            divider.Fill.ForeColor.RGB = accent_color
                            divider.Line.Visible = False
                        except Exception:
                            pass
                    else:
                        tr.Font.Size = 18
                        
        try:
            footer = slide.Shapes.AddTextbox(1, 40, 500, 800, 20)
            footer.TextFrame.TextRange.Text = ""
            footer.TextFrame.TextRange.Font.Name = "JetBrains Mono"
            footer.TextFrame.TextRange.Font.Size = 9
            footer.TextFrame.TextRange.Font.Color.RGB = secondary_color
            footer.TextFrame.TextRange.Font.Italic = True
        except Exception:
            pass
            
    except Exception as e:
        print(f"Error applying theme: {e}")

def _collect_slide_content(topic, titles):
    items = []
    for t in titles:
        q = f"{topic} {t}"
        results = search_google_links(q, 3)
        text = ""
        for _, link in results[:2]:
            s = fetch_page_summary(link, 900)
            if s:
                text += s + "\n\n"
        if not text:
            gen = _gemini_generate_content(topic, t)
            if gen:
                text = gen
            else:
                gen = _groq_generate_content(topic, t)
                if gen:
                    text = gen
                else:
                    text = f"Detailed research on {t}."
        # Clean asterisks/stars from text
        text = text.replace("**", "").replace("*", "")
        items.append(text[:3000])
    return items

def _gemini_generate_content(topic, subtitle):
    try:
        from providers.provider_manager import ProviderManager
        prompt = (
            f"Create concise slide content about '{topic}' focusing on '{subtitle}'. "
            f"Provide 4-7 short bullet points in plain text."
        )
        return ProviderManager.gemini(prompt)
    except Exception:
        return ""

def _groq_generate_content(topic, subtitle):
    try:
        from providers.provider_manager import ProviderManager
        prompt = (
            f"Create concise slide content about '{topic}' focusing on '{subtitle}'. "
            f"Provide 4-7 short bullet points in plain text. Do NOT use markdown stars or asterisks like * or **."
        )
        return ProviderManager.chat(system_prompt=None, messages_history=[], user_prompt=prompt)
    except Exception as e:
        print(f"Error generating content via Groq: {e}")
        return ""

def _groq_generate_outline_titles(topic, count):
    try:
        from providers.provider_manager import ProviderManager
        prompt = f"Create {count} concise slide section titles for a presentation about '{topic}'. Return as a numbered list. Do NOT use markdown stars or asterisks like * or **."
        text = ProviderManager.chat(system_prompt=None, messages_history=[], user_prompt=prompt)
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        titles = []
        for ln in lines:
            # Remove leading numbers/dots
            ln = re.sub(r"^\d+\.\s*", "", ln)
            ln = ln.replace("**", "").replace("*", "")
            if ln and ln.lower() not in [t.lower() for t in titles]:
                titles.append(ln)
        return titles
    except Exception as e:
        print(f"Error generating outline titles via Groq: {e}")
        return []

def _generate_outline_titles(topic, count):
    titles = []
    # Try to derive headings from Google
    results = search_google_links(topic, 6)
    for title, link in results:
        if len(titles) >= count:
            break
        cleaned = title.strip()
        if cleaned and cleaned.lower() not in [t.lower() for t in titles]:
            titles.append(cleaned)
    # If insufficient, use Gemini to create an outline
    if len(titles) < count:
        try:
            from providers.provider_manager import ProviderManager
            text = ProviderManager.gemini(f"Create {count} concise slide section titles for a presentation about '{topic}'. Return as a numbered list.")
            lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
            for ln in lines:
                if len(titles) >= count:
                    break
                # Remove leading numbers/dots
                ln = re.sub(r"^\d+\.\s*", "", ln)
                if ln and ln.lower() not in [t.lower() for t in titles]:
                    titles.append(ln)
        except Exception:
            pass
    # If still insufficient, use Groq
    if len(titles) < count and GROQ_API_KEY:
        groq_titles = _groq_generate_outline_titles(topic, count - len(titles))
        for t in groq_titles:
            if len(titles) >= count:
                break
            if t.lower() not in [x.lower() for x in titles]:
                titles.append(t)
    # Fallback to distinct generic titles
    while len(titles) < count:
        titles.append(f"Slide {len(titles)+1}")
    return titles

def _sanitize_filename(name):
    s = re.sub(r"[\\/:*?\"<>|]", " ", name).strip()
    s = re.sub(r"\s+", " ", s)
    return s if s else "Presentation"

def generate_presentation_content_gemini(topic, slide_count, slide_titles=None):
    """
    Generate structured presentation slide content in a single JSON payload using Gemini.
    """
    try:
        from providers.provider_manager import ProviderManager
        
        titles_str = ""
        if slide_titles:
            titles_str = f" The presentation must cover these slide titles in sequence: {json.dumps(slide_titles)}."
            
        prompt = (
            f"Generate a slide deck presentation content on the topic '{topic}' with exactly {slide_count} slides.{titles_str}\n"
            "Return ONLY a valid JSON object matching the schema below. Do not output any markdown code blocks, preamble, notes, or conversational text. Return only the raw JSON string.\n\n"
            "JSON Schema:\n"
            "{\n"
            "  \"slides\": [\n"
            "    {\n"
            "      \"title\": \"Slide Title\",\n"
            "      \"subtitle\": \"Descriptive tagline (3-7 words)\",\n"
            "      \"layout\": \"bullets\" | \"cards\" | \"split\",\n"
            "      \"content\": {\n"
            "        // If layout is 'bullets':\n"
            "        \"points\": [\"Bullet point 1\", \"Bullet point 2\", \"Bullet point 3\", \"Bullet point 4\"]\n"
            "        // If layout is 'cards':\n"
            "        \"cards\": [\n"
            "          {\"title\": \"Card 1 Title\", \"desc\": \"Card 1 description detail (1-2 sentences)\"},\n"
            "          {\"title\": \"Card 2 Title\", \"desc\": \"Card 2 description detail (1-2 sentences)\"},\n"
            "          {\"title\": \"Card 3 Title\", \"desc\": \"Card 3 description detail (1-2 sentences)\"}\n"
            "        ] // Provide exactly 2 or 3 cards\n"
            "        // If layout is 'split':\n"
            "        \"left_col\": {\"header\": \"Column 1 Header\", \"body\": \"• Detail line 1\\n• Detail line 2\\n• Detail line 3\"},\n"
            "        \"right_col\": {\"header\": \"Column 2 Header\", \"body\": \"• Detail line 1\\n• Detail line 2\\n• Detail line 3\"}\n"
            "      }\n"
            "    }\n"
            "  ]\n"
            "}"
        )
        
        text = ProviderManager.gemini(prompt)
        
        # Clean markdown if present
        if text.startswith("```"):
            text = re.sub(r"^```(?:json)?\s*", "", text)
            text = re.sub(r"\s*```$", "", text)
            
        # Parse JSON
        data = json.loads(text.strip())
        if "slides" in data and isinstance(data["slides"], list):
            slides = data["slides"]
            if len(slides) > 0:
                if len(slides) > slide_count:
                    slides = slides[:slide_count]
                return slides
    except Exception as e:
        print(f"Error generating structured presentation: {e}")
    return None

def generate_presentation_content_groq(topic, slide_count, slide_titles=None):
    """
    Generate structured presentation slide content in a single JSON payload using Groq.
    """
    try:
        from providers.provider_manager import ProviderManager
        
        titles_str = ""
        if slide_titles:
            titles_str = f" The presentation must cover these slide titles in sequence: {json.dumps(slide_titles)}."
            
        prompt = (
            f"Generate a slide deck presentation content on the topic '{topic}' with exactly {slide_count} slides.{titles_str}\n"
            "Return ONLY a valid JSON object matching the schema below. Do not output any markdown code blocks, preamble, notes, or conversational text. Return only the raw JSON string.\n\n"
            "JSON Schema:\n"
            "{\n"
            "  \"slides\": [\n"
            "    {\n"
            "      \"title\": \"Slide Title\",\n"
            "      \"subtitle\": \"Descriptive tagline (3-7 words)\",\n"
            "      \"layout\": \"bullets\" | \"cards\" | \"split\",\n"
            "      \"content\": {\n"
            "        // If layout is 'bullets':\n"
            "        \"points\": [\"Bullet point 1\", \"Bullet point 2\", \"Bullet point 3\", \"Bullet point 4\"]\n"
            "        // If layout is 'cards':\n"
            "        \"cards\": [\n"
            "          {\"title\": \"Card 1 Title\", \"desc\": \"Card 1 description detail (1-2 sentences)\"},\n"
            "          {\"title\": \"Card 2 Title\", \"desc\": \"Card 2 description detail (1-2 sentences)\"},\n"
            "          {\"title\": \"Card 3 Title\", \"desc\": \"Card 3 description detail (1-2 sentences)\"}\n"
            "        ] // Provide exactly 2 or 3 cards\n"
            "        // If layout is 'split':\n"
            "        \"left_col\": {\"header\": \"Column 1 Header\", \"body\": \"• Detail line 1\\n• Detail line 2\\n• Detail line 3\"},\n"
            "        \"right_col\": {\"header\": \"Column 2 Header\", \"body\": \"• Detail line 1\\n• Detail line 2\\n• Detail line 3\"}\n"
            "      }\n"
            "    }\n"
            "  ]\n"
            "}"
        )
        
        text = ProviderManager.chat(
            system_prompt="You are J.A.R.V.I.S., a structured presentation assistant.",
            messages_history=[],
            user_prompt=prompt,
            response_format_json=True
        )
        
        # Clean markdown if present
        if text.startswith("```"):
            text = re.sub(r"^```(?:json)?\s*", "", text)
            text = re.sub(r"\s*```$", "", text)
            
        data = json.loads(text.strip())
        if "slides" in data and isinstance(data["slides"], list):
            slides = data["slides"]
            if len(slides) > 0:
                if len(slides) > slide_count:
                    slides = slides[:slide_count]
                
                # Strip asterisks/stars from text inside slides
                for s in slides:
                    if "title" in s and s["title"]:
                        s["title"] = s["title"].replace("**", "").replace("*", "")
                    if "subtitle" in s and s["subtitle"]:
                        s["subtitle"] = s["subtitle"].replace("**", "").replace("*", "")
                    if "content" in s and s["content"]:
                        content = s["content"]
                        if "points" in content and isinstance(content["points"], list):
                            content["points"] = [p.replace("**", "").replace("*", "") for p in content["points"]]
                        if "cards" in content and isinstance(content["cards"], list):
                            for card in content["cards"]:
                                if "title" in card and card["title"]:
                                    card["title"] = card["title"].replace("**", "").replace("*", "")
                                if "desc" in card and card["desc"]:
                                    card["desc"] = card["desc"].replace("**", "").replace("*", "")
                        if "left_col" in content and isinstance(content["left_col"], dict):
                            if "header" in content["left_col"] and content["left_col"]["header"]:
                                content["left_col"]["header"] = content["left_col"]["header"].replace("**", "").replace("*", "")
                            if "body" in content["left_col"] and content["left_col"]["body"]:
                                content["left_col"]["body"] = content["left_col"]["body"].replace("**", "").replace("*", "")
                        if "right_col" in content and isinstance(content["right_col"], dict):
                            if "header" in content["right_col"] and content["right_col"]["header"]:
                                content["right_col"]["header"] = content["right_col"]["header"].replace("**", "").replace("*", "")
                            if "body" in content["right_col"] and content["right_col"]["body"]:
                                content["right_col"]["body"] = content["right_col"]["body"].replace("**", "").replace("*", "")
                return slides
    except Exception as e:
        print(f"Error generating structured presentation via Groq: {e}")
    return None

def get_fallback_slide_deck(topic, slide_count, slide_titles, contents):
    slides = []
    for i in range(slide_count):
        title = slide_titles[i] if i < len(slide_titles) else f"Slide {i+1}"
        body = contents[i] if i < len(contents) else ""
        
        # Clean paragraphs
        points = []
        for line in body.split("\n"):
            line = line.strip().lstrip("*-#•").strip()
            if line and len(line) > 5:
                # remove preambles
                if not any(f in line.lower() for f in ["here is", "sure, here", "slide content", "focusing on"]):
                    points.append(line)
        if not points:
            points = [body] if body else [f"Detailed research on {title}."]
            
        slides.append({
            "title": title,
            "subtitle": "Overview Report",
            "layout": "bullets",
            "content": {
                "points": points[:5]
            }
        })
    return slides


def _compile_ppt_via_python_pptx(topic, slides_data, theme, out_path):
    """
    Cross-platform .pptx slide deck generator using pure python-pptx.
    Requires zero OS-level PowerPoint or Node.js dependencies.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
    except ImportError:
        return False

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def hex_to_rgb(hex_str):
        hex_str = str(hex_str).lstrip('#')
        if len(hex_str) == 6:
            return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
        return RGBColor(0, 210, 255)

    bg_hex = theme.get('background_gradient', {}).get('start', '#07090D')
    primary_color = hex_to_rgb(theme.get('text', {}).get('primary', '#00D2FF'))
    secondary_color = hex_to_rgb(theme.get('text', {}).get('secondary', '#8DA4B4'))
    accent_color = hex_to_rgb(theme.get('accent', '#00FFCC'))
    bg_color = hex_to_rgb(bg_hex)
    card_bg_color = hex_to_rgb('#0C1522')
    card_border_color = hex_to_rgb('#005F73')

    title_font = theme.get('fonts', {}).get('title', 'Outfit')
    body_font = theme.get('fonts', {}).get('body', 'Outfit')
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = bg_color
        bg.line.fill.background()

    # 1. Cover
    cover = prs.slides.add_slide(blank_layout)
    set_slide_background(cover)

    accent_bar = cover.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.1), Inches(1.8), Inches(0.08))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = accent_color
    accent_bar.line.fill.background()

    tb = cover.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(10.8), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = str(topic).upper()
    p.font.name = title_font
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = primary_color

    p_sub = tf.add_paragraph()
    p_sub.text = 'AUTONOMOUS SYSTEM EXECUTIVE SUMMARY & ARCHITECTURE DECK'
    p_sub.font.name = body_font
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = secondary_color
    p_sub.space_before = Pt(14)

    # 2. Content Slides
    for idx, item in enumerate(slides_data, start=2):
        s = prs.slides.add_slide(blank_layout)
        set_slide_background(s)

        v_bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.7), Inches(0.08), Inches(0.9))
        v_bar.fill.solid()
        v_bar.fill.fore_color.rgb = accent_color
        v_bar.line.fill.background()

        stb = s.shapes.add_textbox(Inches(1.05), Inches(0.6), Inches(11.4), Inches(1.1))
        stf = stb.text_frame
        stf.word_wrap = True
        sp = stf.paragraphs[0]
        sp.text = str(item.get('title', f'Slide {idx}'))
        sp.font.name = title_font
        sp.font.size = Pt(28)
        sp.font.bold = True
        sp.font.color.rgb = primary_color

        if item.get('subtitle'):
            sp2 = stf.add_paragraph()
            sp2.text = str(item.get('subtitle', ''))
            sp2.font.name = body_font
            sp2.font.size = Pt(13)
            sp2.font.color.rgb = secondary_color

        content = item.get('content', {})
        items = []
        if isinstance(content, dict):
            if 'points' in content and isinstance(content['points'], list):
                items = content['points']
            elif 'cards' in content and isinstance(content['cards'], list):
                for c in content['cards']:
                    if isinstance(c, dict):
                        items.append(f"{c.get('heading') or c.get('title') or ''}: {c.get('body') or c.get('desc') or ''}".strip(': '))
                    else:
                        items.append(str(c))
            elif 'stats' in content and isinstance(content['stats'], list):
                for st in content['stats']:
                    if isinstance(st, dict):
                        items.append(f"{st.get('value', '')} — {st.get('label', '')}")
                    else:
                        items.append(str(st))
            elif 'left_col' in content and 'right_col' in content:
                items = (content.get('left_col') or []) + (content.get('right_col') or [])
        elif isinstance(content, list):
            items = content
        elif isinstance(content, str):
            items = [content]

        if not items:
            items = [f'Detailed analysis and strategy for {item.get("title", topic)}.']

        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.7), Inches(4.7))
        card.fill.solid()
        card.fill.fore_color.rgb = card_bg_color
        card.line.color.rgb = card_border_color
        card.line.width = Pt(1)

        ctb = s.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(4.3))
        ctf = ctb.text_frame
        ctf.word_wrap = True

        for p_i, text_line in enumerate(items[:6]):
            cp = ctf.paragraphs[0] if p_i == 0 else ctf.add_paragraph()
            clean_text = str(text_line).strip().lstrip('•-* ').strip()
            cp.text = f'•  {clean_text}'
            cp.font.name = body_font
            cp.font.size = Pt(16)
            cp.font.color.rgb = secondary_color
            cp.space_after = Pt(12)

    # 3. Conclusion Slide
    final_slide = prs.slides.add_slide(blank_layout)
    set_slide_background(final_slide)

    ftb = final_slide.shapes.add_textbox(Inches(1.5), Inches(2.8), Inches(10.3), Inches(2.0))
    ftf = ftb.text_frame
    fp = ftf.paragraphs[0]
    fp.text = 'PRESENTATION COMPLETE'
    fp.alignment = PP_ALIGN.CENTER
    fp.font.name = title_font
    fp.font.size = Pt(40)
    fp.font.bold = True
    fp.font.color.rgb = primary_color

    fp2 = ftf.add_paragraph()
    fp2.text = f'GENERATED BY J.A.R.V.I.S. // TOPIC: {str(topic).upper()}'
    fp2.alignment = PP_ALIGN.CENTER
    fp2.font.name = body_font
    fp2.font.size = Pt(14)
    fp2.font.color.rgb = accent_color
    fp2.space_before = Pt(12)

    prs.save(out_path)
    return True

def _compile_ppt_via_html(topic, slides_data, theme, out_path):
    import shutil
    import subprocess
    import tempfile
    
    background_color = theme.get("background_gradient", {}).get("start", "#07090D")
    text_primary = theme.get("text", {}).get("primary", "#00D2FF")
    text_secondary = theme.get("text", {}).get("secondary", "#8DA4B4")
    accent_color = theme.get("accent", "#00FFCC")
    title_font = theme.get("fonts", {}).get("title", "Outfit")
    body_font = theme.get("fonts", {}).get("body", "Outfit")
    
    import tempfile
    temp_dir = tempfile.mkdtemp(prefix="temp_slides_")
    
    html_files = []
    
    # HTML base template
    html_template = """<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  @import url('https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600;700&family=JetBrains+Mono:ital,wght@0,400;0,700;1,400&display=swap');
  
  body {{
    width: 960px;
    height: 540px;
    margin: 0;
    padding: 50px 60px;
    box-sizing: border-box;
    background-color: {background_color};
    color: {text_secondary};
    font-family: '{body_font}', sans-serif;
    position: relative;
    overflow: hidden;
  }}
  
  h1 {{
    font-family: '{title_font}', sans-serif;
    color: {text_primary};
    font-size: 36px;
    margin: 0 0 5px 0;
    font-weight: 700;
  }}
  
  .subtitle {{
    font-family: 'JetBrains Mono', monospace;
    color: {accent_color};
    font-size: 14px;
    font-weight: 600;
    text-transform: uppercase;
    margin-bottom: 25px;
    letter-spacing: 1px;
  }}
  
  .divider {{
    height: 3px;
    background-color: {accent_color};
    width: 150px;
    margin-bottom: 30px;
    border-radius: 2px;
  }}
  
  .footer {{
    position: absolute;
    bottom: 30px;
    left: 60px;
    font-family: 'JetBrains Mono', monospace;
    font-size: 10px;
    color: {text_secondary};
    opacity: 0.6;
    text-transform: uppercase;
  }}

  /* BULLETS LAYOUT */
  .bullets-container {{
    margin-top: 20px;
  }}
  
  ul {{
    margin: 0;
    padding-left: 20px;
  }}
  
  li {{
    font-size: 18px;
    line-height: 1.6;
    margin-bottom: 12px;
    color: {text_secondary};
  }}

  /* CARDS LAYOUT */
  .cards-container {{
    display: flex;
    gap: 20px;
    margin-top: 10px;
    height: 290px;
  }}
  
  .card {{
    flex: 1;
    background-color: #12171e;
    border: 1px solid {accent_color};
    border-radius: 8px;
    padding: 20px;
    box-sizing: border-box;
    display: flex;
    flex-direction: column;
  }}
  
  .card-title {{
    font-family: '{title_font}', sans-serif;
    font-size: 18px;
    font-weight: 600;
    color: {text_primary};
    margin-bottom: 10px;
  }}
  
  .card-body {{
    font-size: 13px;
    line-height: 1.5;
    color: {text_secondary};
  }}

  /* SPLIT LAYOUT */
  .split-container {{
    display: flex;
    gap: 40px;
    margin-top: 10px;
    height: 290px;
  }}
  
  .column {{
    flex: 1;
    background-color: #12171e;
    border: 1px solid {accent_color};
    border-radius: 8px;
    padding: 20px;
    box-sizing: border-box;
  }}
  
  .col-title {{
    font-family: '{title_font}', sans-serif;
    font-size: 20px;
    font-weight: 600;
    color: {text_primary};
    margin-bottom: 8px;
  }}
  
  .col-divider {{
    height: 1px;
    background-color: {accent_color};
    width: 100%;
    margin-bottom: 12px;
    border-radius: 1px;
  }}
  
  .col-body {{
    font-size: 14px;
    line-height: 1.6;
    color: {text_secondary};
  }}
  
  /* COVER LAYOUT */
  .cover-body {{
    display: flex;
    flex-direction: column;
    justify-content: center;
    align-items: center;
    height: 100%;
    text-align: center;
    padding-top: 50px;
    box-sizing: border-box;
  }}
  
  .cover-title {{
    font-size: 48px;
    font-weight: 700;
    color: {text_primary};
    margin-bottom: 15px;
    text-transform: uppercase;
  }}
  
  .cover-subtitle {{
    font-family: 'JetBrains Mono', monospace;
    font-size: 14px;
    color: {text_secondary};
    letter-spacing: 2px;
    margin-top: 15px;
  }}
  
  .cover-divider {{
    height: 4px;
    background-color: {accent_color};
    width: 250px;
    border-radius: 2px;
  }}
</style>
</head>
<body>
{body_content}
</body>
</html>
"""

    def format_body_text(text):
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        has_bullets = False
        for line in lines:
            if line.startswith(("•", "-", "*")):
                has_bullets = True
                break
        
        if has_bullets:
            points_str = ""
            for line in lines:
                cleaned = line.lstrip("•-* ").strip()
                if cleaned:
                    points_str += f"<li>{cleaned}</li>"
            return f"<ul>{points_str}</ul>"
        else:
            return f"<p>{text.replace(chr(10), '<br>')}</p>"

    def write_html_file(filename, content_html):
        full_html = html_template.format(
            background_color=background_color,
            text_primary=text_primary,
            text_secondary=text_secondary,
            accent_color=accent_color,
            title_font=title_font,
            body_font=body_font,
            body_content=content_html
        )
        file_path = os.path.join(temp_dir, filename)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(full_html)
        html_files.append(file_path)

    # 1. Cover Slide HTML
    cover_html = f"""
    <div class="cover-body">
      <h1 class="cover-title">{topic.upper()}</h1>
      <div class="cover-divider"></div>
      <p class="cover-subtitle">INTELLIGENCE REPORT & SUMMARY OUTLINE</p>
    </div>
    <p class="footer"></p>
    """
    write_html_file("slide_00.html", cover_html)

    # 2. Content Slides HTML
    for s_idx, slide_data in enumerate(slides_data):
        title = slide_data.get("title", f"Section {s_idx + 1}")
        subtitle = slide_data.get("subtitle", "Overview Report")
        layout = slide_data.get("layout", "bullets")
        content = slide_data.get("content", {})
        
        content_html = ""
        if layout == "cards" and "cards" in content and isinstance(content["cards"], list):
            cards_str = ""
            for card in content["cards"][:3]:
                c_title = card.get("title", "")
                c_desc = card.get("desc", "")
                c_desc_formatted = format_body_text(c_desc)
                cards_str += f"""
                <div class="card">
                    <h2 class="card-title">{c_title}</h2>
                    {c_desc_formatted}
                </div>
                """
            content_html = f'<div class="cards-container">{cards_str}</div>'
            
        elif layout == "split" and "left_col" in content and "right_col" in content:
            left = content["left_col"]
            right = content["right_col"]
            left_header = left.get("header", "")
            left_body = left.get("body", "")
            left_body_formatted = format_body_text(left_body)
            right_header = right.get("header", "")
            right_body = right.get("body", "")
            right_body_formatted = format_body_text(right_body)
            
            content_html = f"""
            <div class="split-container">
                <div class="column">
                    <h2 class="col-title">{left_header}</h2>
                    <div class="col-divider"></div>
                    {left_body_formatted}
                </div>
                <div class="column">
                    <h2 class="col-title">{right_header}</h2>
                    <div class="col-divider"></div>
                    {right_body_formatted}
                </div>
            </div>
            """
        else: # bullets
            points_str = ""
            for pt in content.get("points", []):
                points_str += f"<li>{pt}</li>"
            content_html = f'<div class="bullets-container"><ul>{points_str}</ul></div>'

        slide_html = f"""
        <h1>{title}</h1>
        <p class="subtitle">{subtitle.upper()}</p>
        <div class="divider"></div>
        {content_html}
        <p class="footer"></p>
        """
        write_html_file(f"slide_{s_idx+1:02d}.html", slide_html)

    # 3. Conclusion Slide HTML
    conclusion_html = """
    <div class="cover-body">
      <h1 class="cover-title">THANK YOU</h1>
      <div class="cover-divider"></div>
      <p class="cover-subtitle">PRESENTATION COMPLETE</p>
    </div>
    <p class="footer"></p>
    """
    write_html_file("slide_final.html", conclusion_html)

    # 4. Invoke node runner script
    cmd = ["node", "scripts/html2pptx_runner.js", out_path] + html_files
    print(f"Running HTML compiler command: {' '.join(cmd)}")
    
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, check=True)
        print("HTML Slide Compiler Stdout:", res.stdout)
        print("HTML Slide Compiler Stderr:", res.stderr)
        return True
    except subprocess.CalledProcessError as e:
        print("HTML Slide Compiler failed with exit code:", e.returncode)
        print("HTML Slide Compiler Stdout:", e.stdout)
        print("HTML Slide Compiler Stderr:", e.stderr)
        raise e
    finally:
        # Cleanup temp directory
        try:
            shutil.rmtree(temp_dir)
        except Exception as ce:
            print(f"Error deleting temp slide dir: {ce}")

def generate_ppt_custom(topic, slide_count, slide_titles=None, theme=None, existing_ppt_path=None):
    # Adjust slide_count to exclude the cover slide from the requested count
    content_slide_count = max(1, slide_count - 1)
    
    titles = slide_titles[:content_slide_count] if slide_titles else _generate_outline_titles(topic, content_slide_count)
    
    # Try to generate structured presentation payload using Gemini first
    slides_data = generate_presentation_content_gemini(topic, content_slide_count, titles)
    
    # Fallback to Groq structured generation if Gemini fails
    if not slides_data:
        print("Gemini structured generation failed. Trying Groq structured generation...")
        slides_data = generate_presentation_content_groq(topic, content_slide_count, titles)
        
    # Fallback if both structured generations fail
    if not slides_data:
        print("Falling back to standard outlines...")
        contents = _collect_slide_content(topic, titles)
        slides_data = get_fallback_slide_deck(topic, content_slide_count, titles, contents)
        
    base_name = _sanitize_filename(topic)
    out_path = os.path.join(os.getcwd(), f"{base_name}.pptx")
    if os.path.exists(out_path):
        suffix = 1
        while True:
            candidate = os.path.join(os.getcwd(), f"{base_name}_{suffix}.pptx")
            if not os.path.exists(candidate):
                out_path = candidate
                break
            suffix += 1

    if not theme:
        theme = {
            "background_gradient": {"start": "#07090D", "end": "#0A1F30"},
            "text": {"primary": "#00D2FF", "secondary": "#8DA4B4"},
            "fonts": {"title": "Outfit", "body": "Outfit"},
            "accent": "#00FFCC"
        }

    # 1. Try python-pptx (pure Python, cross-platform, Linux/Render compatible)
    try:
        print("Attempting slide generation via python-pptx engine...")
        if _compile_ppt_via_python_pptx(topic, slides_data, theme, out_path):
            print(f"Presentation successfully compiled via python-pptx: {out_path}")
            return out_path
    except Exception as e:
        print(f"python-pptx engine notice: {e}")

    # 2. Try HTML Slide Compiler
    try:
        print("Attempting slide generation via HTML Layout Compiler...")
        compiled = _compile_ppt_via_html(topic, slides_data, theme, out_path)
        if compiled:
            print(f"Presentation successfully compiled via HTML layout engine: {out_path}")
            return out_path
    except Exception as e:
        print(f"HTML Slide Compiler failed ({e}). Checking Windows COM...")

    # 3. Fallback to win32com ONLY if on Windows with Dispatch available
    if win32 is not None and hasattr(win32, 'Dispatch'):
        app = win32.Dispatch("PowerPoint.Application")
    else:
        raise RuntimeError("PowerPoint generation requires python-pptx on Linux/Render or Microsoft PowerPoint on Windows.")
    app.Visible = True
        
    primary_color = _hex_to_rgb(theme["text"]["primary"])
    secondary_color = _hex_to_rgb(theme["text"]["secondary"])
    accent_color = _hex_to_rgb(theme["accent"])
    title_font = theme["fonts"]["title"]
    body_font = theme["fonts"]["body"]
    
    pres = app.Presentations.Add()
    
    # ── SLIDE 1: Cover Slide ───────────────────────────────────────────────
    slide1 = pres.Slides.Add(1, 12) # Blank Layout
    _apply_theme(pres, slide1, theme)
    
    # Main Cover Title (Centered, large)
    title_box = slide1.Shapes.AddTextbox(1, 66, 160, 828, 120)
    title_tr = title_box.TextFrame.TextRange
    title_tr.Text = topic.upper()
    title_tr.ParagraphFormat.Alignment = 2 # Center
    title_tr.Font.Name = title_font
    title_tr.Font.Size = 44
    title_tr.Font.Bold = True
    title_tr.Font.Color.RGB = primary_color
    
    # Add title bottom highlight bar (under cover title)
    cov_bar = slide1.Shapes.AddShape(1, 380, 290, 200, 3) 
    cov_bar.Fill.Solid()
    cov_bar.Fill.ForeColor.RGB = accent_color
    cov_bar.Line.Visible = False
    
    # Cover Subtitle
    sub_box = slide1.Shapes.AddTextbox(1, 66, 310, 828, 60)
    sub_tr = sub_box.TextFrame.TextRange
    sub_tr.Text = "INTELLIGENCE REPORT & SUMMARY OUTLINE"
    sub_tr.ParagraphFormat.Alignment = 2 # Center
    sub_tr.Font.Name = "JetBrains Mono"
    sub_tr.Font.Size = 11
    sub_tr.Font.Color.RGB = secondary_color
    
    # ── SLIDES 2 to N: Custom Layout Slides ──────────────────────────────────
    idx = 2
    for s_idx, slide_data in enumerate(slides_data):
        slide = pres.Slides.Add(idx, 12) # Blank Layout
        idx += 1
        
        # Apply theme (sets background and footer)
        _apply_theme(pres, slide, theme)
        
        # Add slide title textbox
        title_box = slide.Shapes.AddTextbox(1, 66, 28.8, 828, 60)
        title_box.Name = "SlideTitle"
        t_tr = title_box.TextFrame.TextRange
        t_tr.Text = slide_data.get("title", f"Section {s_idx + 1}")
        t_tr.Font.Name = title_font
        t_tr.Font.Size = 32
        t_tr.Font.Bold = True
        t_tr.Font.Color.RGB = primary_color
        
        # Add thin divider line under title
        div_line = slide.Shapes.AddShape(1, 66, 95, 828, 1.5)
        div_line.Fill.Solid()
        div_line.Fill.ForeColor.RGB = accent_color
        div_line.Line.Visible = False
        
        # Optional Subtitle/Tagline
        sub_text = slide_data.get("subtitle", "")
        if sub_text:
            tag_box = slide.Shapes.AddTextbox(1, 66, 102, 828, 24)
            tag_box.Name = "SlideSubtitle"
            tag_tr = tag_box.TextFrame.TextRange
            tag_tr.Text = sub_text.upper()
            tag_tr.Font.Name = "JetBrains Mono"
            tag_tr.Font.Size = 9
            tag_tr.Font.Bold = True
            tag_tr.Font.Color.RGB = secondary_color
            
        # Draw Content based on Layout
        layout = slide_data.get("layout", "bullets")
        content = slide_data.get("content", {})
        
        if layout == "cards" and "cards" in content and isinstance(content["cards"], list):
            cards = content["cards"][:3] # Maximum 3 cards
            num_cards = len(cards)
            
            gap = 24
            top_pos = 155
            height_pos = 295
            
            if num_cards == 3:
                card_width = 260
                positions = [66, 66 + 260 + gap, 66 + 2 * (260 + gap)]
            elif num_cards == 2:
                card_width = 402
                positions = [66, 66 + 402 + gap]
            else:
                card_width = 828
                positions = [66]
                
            for c_i, card_data in enumerate(cards):
                left_pos = positions[c_i]
                
                # Card shape
                card = slide.Shapes.AddShape(5, left_pos, top_pos, card_width, height_pos)
                card.Name = f"CardShape_{c_i}"
                card.Fill.Solid()
                card.Fill.ForeColor.RGB = 0x1A140F 
                
                card.Line.Visible = True
                card.Line.ForeColor.RGB = accent_color
                card.Line.Weight = 1.5
                
                tf = card.TextFrame
                tf.MarginLeft = 14
                tf.MarginRight = 14
                tf.MarginTop = 14
                tf.MarginBottom = 14
                
                c_title = card_data.get("title", "")
                c_desc = card_data.get("desc", "")
                
                tr = tf.TextRange
                tr.Text = f"{c_title}\n\n{c_desc}"
                
                # Format card title
                p1 = tr.Paragraphs(1)
                p1.Font.Name = title_font
                p1.Font.Size = 15
                p1.Font.Bold = True
                p1.Font.Color.RGB = primary_color
                
                # Format card body
                if tr.Paragraphs().Count >= 3:
                    for pi in range(3, tr.Paragraphs().Count + 1):
                        p = tr.Paragraphs(pi)
                        p.Font.Name = body_font
                        p.Font.Size = 11.5
                        p.Font.Bold = False
                        p.Font.Color.RGB = secondary_color
                        
        elif layout == "split" and "left_col" in content and "right_col" in content:
            left_col = content["left_col"]
            right_col = content["right_col"]
            
            col_width = 394
            col_gap = 40
            col_top = 155
            col_height = 295
            
            cols = [left_col, right_col]
            for c_i, col_data in enumerate(cols):
                left_pos = 66 + c_i * (col_width + col_gap)
                
                # Column shape
                col_shape = slide.Shapes.AddShape(5, left_pos, col_top, col_width, col_height)
                col_shape.Name = f"ColShape_{c_i}"
                col_shape.Fill.Solid()
                col_shape.Fill.ForeColor.RGB = 0x1A140F
                
                col_shape.Line.Visible = True
                col_shape.Line.ForeColor.RGB = accent_color
                col_shape.Line.Weight = 1.5
                
                tf = col_shape.TextFrame
                tf.MarginLeft = 14
                tf.MarginRight = 14
                tf.MarginTop = 14
                tf.MarginBottom = 14
                
                col_header = col_data.get("header", "")
                col_body = col_data.get("body", "")
                
                tr = tf.TextRange
                tr.Text = f"{col_header}\n\n{col_body}"
                
                # Format header
                p1 = tr.Paragraphs(1)
                p1.Font.Name = title_font
                p1.Font.Size = 17
                p1.Font.Bold = True
                p1.Font.Color.RGB = primary_color
                
                # Format body
                if tr.Paragraphs().Count >= 3:
                    for pi in range(3, tr.Paragraphs().Count + 1):
                        p = tr.Paragraphs(pi)
                        p.Font.Name = body_font
                        p.Font.Size = 12.5
                        p.Font.Bold = False
                        p.Font.Color.RGB = secondary_color
                        
        else: # Standard Bullets layout
            points = content.get("points", [])
            if not points:
                points = [f"Detailed analysis on {slide_data.get('title') or 'topic'}."]
                
            # Add content textbox
            bbox = slide.Shapes.AddTextbox(1, 66, 155, 828, 300)
            bbox.Name = "BulletBox"
            tr = bbox.TextFrame.TextRange
            
            # Format text as a bulleted list
            bullets_text = "\n".join([f"•   {pt}" for pt in points])
            tr.Text = bullets_text
            
            for pi in range(1, tr.Paragraphs().Count + 1):
                p = tr.Paragraphs(pi)
                p.Font.Name = body_font
                p.Font.Size = 17
                p.Font.Color.RGB = secondary_color
                p.Font.Bold = False
                
    # ── SLIDE N+1: Conclusion Slide ─────────────────────────────────────────
    thank_slide = pres.Slides.Add(idx, 12)
    _apply_theme(pres, thank_slide, theme)
    
    thank_box = thank_slide.Shapes.AddTextbox(1, 66, 200, 828, 100)
    thank_tr = thank_box.TextFrame.TextRange
    thank_tr.Text = "THANK YOU"
    thank_tr.ParagraphFormat.Alignment = 2 # Center
    thank_tr.Font.Name = title_font
    thank_tr.Font.Size = 48
    thank_tr.Font.Bold = True
    thank_tr.Font.Color.RGB = primary_color
    
    thank_bar = thank_slide.Shapes.AddShape(1, 430, 310, 100, 3)
    thank_bar.Fill.Solid()
    thank_bar.Fill.ForeColor.RGB = accent_color
    thank_bar.Line.Visible = False
    
    pres.SaveAs(out_path)
    return out_path

def run_ppt_gui():
    root = tk.Tk()
    root.title("JARVIS — Generate PPT")
    root.geometry("600x520")
    style = ttk.Style()
    style.theme_use('clam')

    frm = ttk.Frame(root)
    frm.pack(fill=tk.BOTH, expand=True, padx=12, pady=12)

    ttk.Label(frm, text="Topic").pack(anchor=tk.W)
    topic_entry = ttk.Entry(frm, width=68)
    topic_entry.pack(fill=tk.X, pady=(2, 8))

    ttk.Label(frm, text="Number of slides").pack(anchor=tk.W)
    slides_entry = ttk.Entry(frm, width=12)
    slides_entry.insert(0, "6")
    slides_entry.pack(pady=(2, 8))

    ttk.Label(frm, text="Custom slide format?").pack(anchor=tk.W)
    format_var = tk.StringVar(value="no")
    fmt_row = ttk.Frame(frm)
    fmt_row.pack(pady=(2, 8))
    ttk.Radiobutton(fmt_row, text="Yes", variable=format_var, value="yes").pack(side=tk.LEFT, padx=6)
    ttk.Radiobutton(fmt_row, text="No", variable=format_var, value="no").pack(side=tk.LEFT, padx=6)

    titles_container = ttk.Frame(frm)
    titles_container.pack(fill=tk.X, pady=(4, 8))
    title_entries = []

    def rebuild_titles():
        for w in titles_container.winfo_children():
            w.destroy()
        title_entries.clear()
        if format_var.get() != "yes":
            return
        try:
            cnt = max(1, int(slides_entry.get()))
        except Exception:
            cnt = 6
        for i in range(cnt):
            row = ttk.Frame(titles_container)
            row.pack(fill=tk.X, pady=2)
            ttk.Label(row, text=f"Slide {i+1}").pack(side=tk.LEFT, padx=6)
            e = ttk.Entry(row, width=50)
            e.pack(side=tk.LEFT)
            title_entries.append(e)

    format_var.trace_add('write', lambda *args: rebuild_titles())
    slides_entry.bind('<FocusOut>', lambda e: rebuild_titles())

    ttk.Separator(frm).pack(fill=tk.X, pady=6)
    ttk.Label(frm, text="Existing themed PPT to reuse?").pack(anchor=tk.W)
    reuse_var = tk.StringVar(value="no")
    reuse_row = ttk.Frame(frm)
    reuse_row.pack(pady=(2, 8))
    ttk.Radiobutton(reuse_row, text="Yes", variable=reuse_var, value="yes").pack(side=tk.LEFT, padx=6)
    ttk.Radiobutton(reuse_row, text="No", variable=reuse_var, value="no").pack(side=tk.LEFT, padx=6)
    existing_path_var = tk.StringVar()
    sel_row = ttk.Frame(frm)
    sel_row.pack(fill=tk.X)
    ttk.Entry(sel_row, textvariable=existing_path_var, width=50).pack(side=tk.LEFT, padx=6)
    ttk.Button(sel_row, text="Browse", command=lambda: existing_path_var.set(select_existing_ppt() or "")).pack(side=tk.LEFT)

    ttk.Separator(frm).pack(fill=tk.X, pady=6)
    ttk.Label(frm, text="Use Gemini for theme?").pack(anchor=tk.W)
    use_gen_var = tk.StringVar(value="yes")
    gen_row = ttk.Frame(frm)
    gen_row.pack(pady=(2, 8))
    ttk.Radiobutton(gen_row, text="Yes", variable=use_gen_var, value="yes").pack(side=tk.LEFT, padx=6)
    ttk.Radiobutton(gen_row, text="No", variable=use_gen_var, value="no").pack(side=tk.LEFT, padx=6)

    status = ttk.Label(frm, text="")
    status.pack(pady=8)

    result = {"path": None}

    def on_generate():
        topic = topic_entry.get().strip()
        if not topic:
            status.config(text="Enter a topic")
            return
        status.config(text="Generating presentation... Please wait")
        generate_btn.configure(state=tk.DISABLED)
        try:
            slide_count = max(1, int(slides_entry.get()))
        except Exception:
            slide_count = 6
        if format_var.get() == "yes":
            titles = [e.get().strip() or f"Slide {i+1}" for i, e in enumerate(title_entries)]
            if len(titles) != slide_count:
                status.config(text="Provide titles for all slides")
                generate_btn.configure(state=tk.NORMAL)
                return
        else:
            titles = _generate_outline_titles(topic, slide_count)
        theme = None
        existing = existing_path_var.get().strip() if reuse_var.get() == "yes" else None
        try:
            path = generate_ppt_custom(topic, slide_count, slide_titles=titles, theme=theme, existing_ppt_path=existing)
            status.config(text=f"Created {os.path.abspath(path)}")
            result["path"] = path
            root.after(800, root.destroy)
        except Exception as e:
            status.config(text=f"Error: {e}")
            generate_btn.configure(state=tk.NORMAL)

    generate_btn = ttk.Button(frm, text="Generate", command=on_generate)
    generate_btn.pack(pady=10)

    root.mainloop()
    return result["path"]